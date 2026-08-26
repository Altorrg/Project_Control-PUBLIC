from __future__ import annotations

from datetime import date
from decimal import Decimal
import unittest
from pydantic import ValidationError

from alkalmazas.adatmodellek.penzugy import PaymentCreate
from alkalmazas.szolgaltatasok.engedelyezes import latest_round_date, weighted_progress
from alkalmazas.szolgaltatasok.alvallalkozok import (
    build_subcontractor_workbook,
    matrix_payment_reference,
    normalize_workflow_money,
    offer_part_total,
    validate_step_close,
)
from alkalmazas.szolgaltatasok.penzugy import (
    build_finance_items,
    build_project_finance_workbook,
    company_cash_flow,
    finance_allocations,
    finance_summary,
    labor_trends,
    parse_money,
    reconciliation_report,
)
from alkalmazas.szolgaltatasok.utemterv import (
    build_schedule_workbook,
    calculate_milestone_dates,
    calculate_schedule,
    hungarian_holidays,
    is_workday,
)


class ScheduleServiceTests(unittest.TestCase):
    def test_project_start_moves_from_weekend_and_dependencies_follow(self):
        rows = [
            {"id": 1, "position": 10, "phase_name": "Elso", "reference_name": "Projekt kezdete", "days": 2, "delay": 0, "calc_type": "Munkanap"},
            {"id": 2, "position": 20, "phase_name": "Masodik", "reference_name": "Elso", "days": 1, "delay": 0, "calc_type": "Naptári nap"},
        ]
        result = calculate_schedule(rows, date(2026, 7, 18))
        self.assertEqual(result[0]["start_date"], date(2026, 7, 20))
        self.assertEqual(result[0]["end_date"], date(2026, 7, 21))
        self.assertEqual(result[1]["start_date"], date(2026, 7, 22))

    def test_calendar_exceptions_override_weekends(self):
        saturday = date(2026, 7, 18)
        self.assertFalse(is_workday(saturday))
        self.assertTrue(is_workday(saturday, extra_workdays={saturday}))
        self.assertIn(date(2026, 4, 3), hungarian_holidays(2026))

    def test_circular_reference_is_rejected(self):
        rows = [
            {"id": 1, "position": 1, "phase_name": "A", "reference_name": "B", "days": 1},
            {"id": 2, "position": 2, "phase_name": "B", "reference_name": "A", "days": 1},
        ]
        with self.assertRaisesRegex(ValueError, "Körkörös"):
            calculate_schedule(rows, date(2026, 1, 5))

    def test_dynamic_milestones_follow_phase_end_and_each_other(self):
        schedule = calculate_schedule(
            [{"id": 1, "position": 10, "phase_name": "Tervezés", "reference_name": "Projekt kezdete", "days": 5}],
            date(2026, 7, 20),
        )
        result = calculate_milestone_dates(
            [
                {"id": 1, "name": "E01", "date_type": "Ütemterv sorhoz kötve", "reference_name": "Tervezés", "offset_days": 2},
                {"id": 2, "name": "E02", "date_type": "Másik mérföldkőhöz kötve", "reference_name": "E01", "offset_days": 3},
            ],
            schedule,
        )
        self.assertEqual(result[0]["planned_date"], date(2026, 7, 26))
        self.assertEqual(result[1]["planned_date"], date(2026, 7, 29))

    def test_milestone_cycle_is_rejected(self):
        milestones = [
            {"name": "A", "date_type": "Másik mérföldkőhöz kötve", "reference_name": "B"},
            {"name": "B", "date_type": "Másik mérföldkőhöz kötve", "reference_name": "A"},
        ]
        with self.assertRaisesRegex(ValueError, "Körkörös"):
            calculate_milestone_dates(milestones, [])

    def test_schedule_export_contains_timeline_and_milestone(self):
        rows = calculate_schedule(
            [{"id": 1, "position": 10, "phase_name": "Tervezés", "reference_name": "Projekt kezdete", "days": 2}],
            date(2026, 7, 20),
        )
        output = build_schedule_workbook(
            "Minta projekt",
            date(2026, 7, 20),
            rows,
            [{"id": 4, "name": "E01", "phase_no": 1, "sequence_no": 1, "planned_date": date(2026, 7, 22)}],
        )
        self.assertGreater(len(output.getvalue()), 1000)


class EngagementServiceTests(unittest.TestCase):
    def test_weighted_progress_uses_closed_weight(self):
        result = weighted_progress([
            {"weight": 1, "is_closed": True},
            {"weight": 3, "is_closed": False},
        ])
        self.assertEqual(result["progress_percent"], Decimal("25.00"))
        self.assertEqual(result["open_count"], 1)

    def test_latest_round_date_considers_sent_and_received(self):
        result = latest_round_date([
            {"sent_date": "2026-01-01", "received_date": "2026-01-05"},
            {"sent_date": "2026-02-01", "received_date": None},
        ])
        self.assertEqual(result, date(2026, 2, 1))


class SubcontractorServiceTests(unittest.TestCase):
    def test_step_one_requires_company_and_discipline(self):
        with self.assertRaisesRegex(ValueError, "kötelező"):
            validate_step_close({"company_name": "", "discipline": "Útépítés"}, [], 1)

    def test_offer_parts_ignore_optional_rows_and_require_matching_total(self):
        parts = [
            {"phase_name": "Alap", "amount": "1 000 000", "optional": False},
            {"phase_name": "Opció", "amount": "250 000", "optional": True},
        ]
        self.assertEqual(offer_part_total(parts), Decimal("1000000"))
        with self.assertRaisesRegex(ValueError, "nem egyezik"):
            validate_step_close({"offer_price": "900 000", "no_offer": False}, parts, 3)
        warnings = validate_step_close(
            {"offer_price": "900 000", "no_offer": False},
            parts,
            3,
            allow_offer_mismatch=True,
        )
        self.assertEqual(len(warnings), 1)

    def test_accepted_contract_requires_matrix_base_and_normalizes_total(self):
        with self.assertRaisesRegex(ValueError, "mátrix"):
            validate_step_close(
                {"acceptance": "Elfogadva", "contract_fee": "", "copyright_fee": ""},
                [],
                5,
            )
        normalized = normalize_workflow_money({
            "contract_fee": "8 000 000 Ft",
            "copyright_fee": "2 000 000 Ft",
            "offer_price": "10 000 000",
            "no_offer": False,
        })
        self.assertEqual(normalized["contract_total"], "10000000.00")

    def test_matrix_payment_reference_is_stable_per_workflow_and_column(self):
        self.assertEqual(
            matrix_payment_reference("ALV_7_SZ2_MD", "M_12"),
            "KIA_WF_7_COL_M_12",
        )
        self.assertIsNone(matrix_payment_reference("FO_SZ1_MD", "M_12"))

    def test_subcontractor_export_preserves_folder_order_and_offer_details(self):
        output = build_subcontractor_workbook(
            [{"id": 2, "name": "Környezetvédelem", "display_order": 10}],
            [{
                "id": 7,
                "discipline": "Környezetvédelem",
                "company_name": "Minta Kft",
                "code": "E1-EX",
                "acceptance": "Elfogadva",
                "signature_status": "Aláírva",
                "offer_parts": [{"phase_name": "ENG", "amount": "2000000", "optional": False}],
            }],
        )
        self.assertGreater(len(output.getvalue()), 1000)
        from openpyxl import load_workbook
        workbook = load_workbook(output, read_only=False)
        sheet = workbook["Ajánlatkérések"]
        self.assertEqual(sheet["A2"].value, "Környezetvédelem")
        self.assertEqual(sheet["C3"].value, "Minta Kft")
        self.assertIn("ENG", sheet["L3"].value)


class FinanceServiceTests(unittest.TestCase):
    def test_payment_status_requires_its_business_dates(self):
        with self.assertRaises(ValidationError):
            PaymentCreate(project_id=1, amount=100, status="Számlázva", invoice_number="SZ-1")
        completed = PaymentCreate(
            project_id=1,
            amount=100,
            status="Pénzügyileg teljesítve",
            tig_date=date(2026, 7, 1),
            invoice_date=date(2026, 7, 2),
            invoice_number="SZ-1",
            payment_date=date(2026, 7, 10),
        )
        self.assertEqual(completed.status, "Pénzügyileg teljesítve")

    def test_milestones_and_matrix_build_linked_finance_items(self):
        items = build_finance_items(
            [{"section_no": 1, "original_fee": 1_000_000, "modified_fee": 0, "fee_percent": 90, "copyright_percent": 10}],
            [{"id": 12, "name": "E01", "phase_no": 1, "planned_date": "2026-08-01", "fee_percent": 20, "copyright_percent": 50}],
            [{"id": 7, "company_name": "Minta Kft", "acceptance": "Elfogadva", "is_deleted": False}],
            [
                {"reference_id": "ALV_7_SZ1_MD", "column_code": "M_12", "amount": 100_000},
                {"reference_id": "ALV_7_SZ1_SZJD", "column_code": "M_12", "amount": 50_000},
            ],
            [],
        )
        revenue = next(item for item in items if item["is_revenue"])
        cost = next(item for item in items if not item["is_revenue"])
        self.assertEqual(revenue["reference"], "BEV_MS_12")
        self.assertEqual(revenue["expected_amount"], Decimal("230000"))
        self.assertEqual(cost["reference"], "KIA_WF_7_COL_M_12")
        self.assertEqual(cost["expected_amount"], Decimal("150000"))

    def test_inactive_workflow_matrix_cells_do_not_create_finance_items(self):
        items = build_finance_items(
            [],
            [{"id": 12, "name": "E01", "phase_no": 1, "planned_date": "2026-08-01"}],
            [{"id": 7, "company_name": "Régi Kft", "acceptance": "Elutasítva", "is_deleted": False}],
            [{"reference_id": "ALV_7_SZ1_MD", "column_code": "M_12", "amount": 100_000}],
            [],
        )
        self.assertEqual(items, [])

    def test_extra_matrix_stage_creates_linked_finance_item(self):
        items = build_finance_items(
            [],
            [],
            [{"id": 7, "company_name": "Minta Kft", "acceptance": "Elfogadva", "is_deleted": False}],
            [{"reference_id": "ALV_7_SZ1_MD", "column_code": "STAGE_4", "amount": 125_000}],
            [],
            [{"id": 4, "name": "Első előleg", "stage_date": "2026-09-01"}],
        )
        self.assertEqual(items[0]["reference"], "KIA_WF_7_COL_STAGE_4")
        self.assertEqual(items[0]["expected_date"], "2026-09-01")
        self.assertEqual(items[0]["expected_amount"], Decimal("125000"))

    def test_summary_preserves_desktop_estimation(self):
        result = finance_summary(
            [{"original_fee": 1_000_000, "modified_fee": 0}],
            0,
            {2026: 100_000},
            [{"discipline_name": "Ut", "record_date": "2026-01-01", "days": 2, "actual_percent": 50}],
            [{"id": 7, "acceptance": "Elfogadva", "contract_total": "200 000 Ft"}],
            [{"workflow_id": 7, "amount": 50_000, "is_revenue": False, "status": "Pénzügyileg teljesítve"}],
        )
        self.assertEqual(result["estimated_internal"], Decimal("400000"))
        self.assertEqual(result["external_contracted"], Decimal("200000"))
        self.assertEqual(result["estimated_profit"], Decimal("400000"))
        self.assertEqual(parse_money("1 250 000 Ft"), Decimal("1250000"))
        self.assertEqual(parse_money("1.250.000,50 Ft"), Decimal("1250000.50"))

    def test_summary_estimates_internal_cost_after_discipline_aggregation(self):
        result = finance_summary(
            [],
            2_000_000,
            {2026: 100_000},
            [
                {"discipline_name": "Útépítés", "record_date": "2026-01-01", "days": 1, "actual_percent": 10},
                {"discipline_name": "Útépítés", "record_date": "2026-02-01", "days": 1, "actual_percent": 10},
                {"discipline_name": "Vízépítés", "record_date": "2026-01-01", "days": 3, "actual_percent": 50},
            ],
            [],
            [],
        )

        self.assertEqual(result["by_discipline"]["Útépítés"]["current"], Decimal("200000"))
        self.assertEqual(result["by_discipline"]["Útépítés"]["estimated"], Decimal("1000000"))
        self.assertEqual(result["by_discipline"]["Vízépítés"]["estimated"], Decimal("600000"))
        self.assertEqual(result["estimated_internal"], Decimal("1600000"))
        self.assertEqual(
            set(result["by_discipline"]["Útépítés"]),
            {"current", "estimated"},
        )

    def test_summary_zero_progress_uses_current_discipline_cost(self):
        result = finance_summary(
            [],
            1_000_000,
            {2026: 100_000},
            [
                {"discipline_name": "Hídépítés", "record_date": "2026-01-01", "days": 2, "actual_percent": 0},
                {"discipline_name": "Hídépítés", "record_date": "2026-02-01", "days": 1, "actual_percent": 0},
            ],
            [],
            [],
        )

        self.assertEqual(result["current_internal"], Decimal("300000"))
        self.assertEqual(result["estimated_internal"], Decimal("300000"))
        self.assertEqual(result["by_discipline"]["Hídépítés"]["estimated"], Decimal("300000"))

    def test_summary_uses_latest_contract_modification(self):
        result = finance_summary(
            [],
            1_000_000,
            {},
            [],
            [{"id": 7, "acceptance": "Elfogadva", "contract_total": "200000", "modified_total": "240000"}],
            [],
        )
        self.assertEqual(result["external_contracted"], Decimal("240000"))
        self.assertEqual(result["estimated_profit"], Decimal("760000"))

    def test_orphaned_derived_payment_is_excluded_from_summary(self):
        result = finance_summary(
            [],
            1_000_000,
            {},
            [],
            [{"id": 7, "acceptance": "Elfogadva", "contract_total": "200000"}],
            [{"workflow_id": 7, "reference": "KIA_WF_7_COL_M_999", "amount": 50_000, "is_revenue": False, "status": "Pénzügyileg teljesítve"}],
            set(),
        )
        self.assertEqual(result["external_paid"], Decimal("0"))
        self.assertEqual(result["subcontractors"][0]["remaining"], Decimal("200000"))

    def test_company_cash_flow_matches_monthly_desktop_logic(self):
        result = company_cash_flow(
            2026,
            [{"id": 1, "name": "Teszt", "contract_value": 1_000_000}],
            [
                {"project_id": 1, "payment_date": "2026-01-10", "amount": 300_000, "is_revenue": True},
                {"project_id": 1, "payment_date": "2026-01-20", "amount": 50_000, "is_revenue": False},
            ],
            [{"project_id": 1, "year": 2026, "daily_rate": 100_000}],
            [{"project_id": 1, "record_date": "2026-01-01", "days": 2}],
        )
        self.assertEqual(result["monthly_revenue"][0], Decimal("300000"))
        self.assertEqual(result["monthly_cost"][0], Decimal("250000"))
        self.assertEqual(result["profit_total"], Decimal("50000"))
        self.assertEqual(result["projects"][0]["annual_net"], Decimal("50000"))

    def test_labor_trends_return_monthly_and_cumulative_days(self):
        result = labor_trends(
            [
                {"discipline_name": "Út", "record_date": "2026-01-01", "days": 2, "planned_percent": 10, "actual_percent": 8},
                {"discipline_name": "Út", "record_date": "2026-02-01", "days": 3, "planned_percent": 15, "actual_percent": 12},
            ],
            {2026: 100_000},
        )
        self.assertEqual(result["months"], ["2026-01", "2026-02"])
        self.assertEqual(result["series"][0]["monthly_days"], [Decimal("2"), Decimal("3")])
        self.assertEqual(result["series"][0]["cumulative_days"], [Decimal("2"), Decimal("5")])

    def test_finance_allocations_keep_unassigned_subcontractors_visible(self):
        result = finance_allocations(
            {"Út": {"current": 100, "estimated": 200}},
            [
                {"id": 7, "company_name": "Minta Kft", "contracted": 300, "paid": 50},
                {"id": 8, "company_name": "Másik Kft", "contracted": 100, "paid": 0},
            ],
            [{"workflow_id": 7, "discipline_name": "Út"}],
            ["Út"],
        )
        self.assertEqual(result["disciplines"][0]["external_contracted"], Decimal("300"))
        self.assertEqual(result["unassigned"]["external_contracted"], Decimal("100"))

    def test_reconciliation_reports_orphans_duplicates_and_mismatch(self):
        report = reconciliation_report(
            [{"reference": "BEV_MS_1", "expected_amount": 100, "expected_date": None}],
            [
                {"id": 1, "reference": "BEV_MS_1", "amount": 90, "status": "Pénzügyileg teljesítve"},
                {"id": 2, "reference": "BEV_MS_1", "amount": 100, "status": "Várható"},
                {"id": 3, "reference": "KIA_WF_7_COL_M_99", "amount": 50, "status": "Várható"},
            ],
        )
        self.assertEqual(report["error_count"], 2)
        self.assertEqual(report["warning_count"], 2)

    def test_project_finance_export_contains_all_operational_sheets(self):
        output = build_project_finance_workbook("Minta", {
            "budget": 1_000_000,
            "estimated_internal": 100_000,
            "external_contracted": 200_000,
            "estimated_profit": 700_000,
            "allocations": {"disciplines": [], "unassigned": {}},
            "labor_trends": {"months": [], "series": []},
            "finance_items": [],
        })
        from openpyxl import load_workbook
        workbook = load_workbook(output)
        self.assertEqual(workbook.sheetnames, ["Vezetői dashboard", "Költségallokáció", "Munkaórák", "Pénzügyi folyamatok"])


if __name__ == "__main__":
    unittest.main()
