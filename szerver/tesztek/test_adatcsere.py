from __future__ import annotations

from io import BytesIO
import importlib.util
import unittest

from openpyxl import Workbook, load_workbook

from alkalmazas.szolgaltatasok.adatcsere import build_engagement_excel, build_engagement_pdf, parse_engagement_workbook


class ImportExportTests(unittest.TestCase):
    def workbook_bytes(self) -> bytes:
        workbook = Workbook()
        data = workbook.active
        data.title = "Adatlap"
        data.append(["Feladat", "Súly", "Aktív", "Sorrend", "Leírás"])
        data.append(["útépítési engedély", 2, True, 10, "Teszt"])
        names = workbook.create_sheet("Névadatlap")
        names.append(["Név", "Cím", "Ügyintéző", "E-mail", "Telefonszám", "Szervezet típus"])
        names.append(["Árvíztűrő Partner", "Őr utca 1.", "Ügyintéző", "partner@example.com", "+36 1", "Hatóság"])
        discipline = workbook.create_sheet("Útépítés")
        discipline.append(["Név", "Cím", "Feladat", "Súly", "Megjegyzés", "Utolsó dátum", "Státusz", "Érvényes", "1. KI dátum", "1. KI megjegyzés", "1. BE dátum", "1. BE megjegyzés"])
        discipline.append(["Árvíztűrő Partner", "Őr utca 1.", "útépítési engedély", 2, "Megjegyzés", "2026-01-10", "○", "2027-01-10", "2026-01-01", "Kiküldve", "2026-01-10", "Beérkezett"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def test_import_parser_keeps_hungarian_text_and_rounds(self):
        parsed = parse_engagement_workbook(self.workbook_bytes(), "minta.xlsx")
        self.assertEqual(parsed["summary"]["assignments"], 1)
        self.assertEqual(parsed["summary"]["rounds"], 1)
        self.assertEqual(parsed["partners"][0]["name"], "Árvíztűrő Partner")
        self.assertEqual(parsed["disciplines"][0]["assignments"][0]["status_code"], "○")

    def test_excel_export_has_compatible_sheet_structure(self):
        data = {
            "tasks": [{"name": "útépítési engedély", "weight": 1, "active": True, "sort_order": 1, "description": ""}],
            "partners": [{"name": "Partner", "address": "", "contact_person": "", "email": "", "phone": "", "org_type": ""}],
            "dashboard": {"disciplines": []},
            "disciplines": [{"name": "Útépítés", "assignments": []}],
        }
        content = build_engagement_excel(data)
        workbook = load_workbook(BytesIO(content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Adatlap", "Névadatlap", "Összegzés", "Útépítés"])
        self.assertEqual(workbook["Útépítés"].max_column, 48)
        workbook.close()

    def test_pdf_export_embeds_unicode_capable_font(self):
        if importlib.util.find_spec("reportlab") is None:
            self.skipTest("A reportlab nincs telepitve ebben a tesztkornyezetben.")
        content = build_engagement_pdf({"dashboard": {"total": 1, "closed": 0, "attention": 1, "disciplines": []}})
        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 1000)


if __name__ == "__main__":
    unittest.main()
