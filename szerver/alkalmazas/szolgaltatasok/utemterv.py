from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Iterable, Mapping


PROJECT_START_LABEL = "Projekt kezdete"
FIXED_DATE = "Fix dátum"
SCHEDULE_LINKED_DATE = "Ütemterv sorhoz kötve"
MILESTONE_LINKED_DATE = "Másik mérföldkőhöz kötve"
MILESTONE_DATE_TYPES = {FIXED_DATE, SCHEDULE_LINKED_DATE, MILESTONE_LINKED_DATE}


def _as_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _as_positive_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return None
    return parsed if parsed > 0 else None


def normalize_milestone(
    item: Mapping[str, Any],
    schedule_rows: Iterable[Mapping[str, Any]] | None = None,
) -> dict[str, Any]:
    normalized = dict(item)
    legacy_name = str(normalized.get("name") or "").strip()
    official_name = str(normalized.get("official_name") or "").strip() or legacy_name
    short_name = str(normalized.get("short_name") or "").strip() or legacy_name or official_name
    official_name = official_name or short_name
    normalized["name"] = official_name
    normalized["short_name"] = short_name
    normalized["official_name"] = official_name
    normalized["description"] = str(normalized.get("description") or "")
    schedule_row_id = _as_positive_int(normalized.get("schedule_row_id"))
    normalized["schedule_row_id"] = schedule_row_id
    if schedule_rows is not None and schedule_row_id is not None:
        linked_row = next(
            (
                row
                for row in schedule_rows
                if _as_positive_int(row.get("id")) == schedule_row_id
            ),
            None,
        )
        if linked_row is None:
            normalized["schedule_row_id"] = None
        else:
            normalized["reference_name"] = str(linked_row.get("phase_name") or "").strip()
    return normalized


def hungarian_holidays(year: int) -> set[date]:
    fixed = {
        date(year, 1, 1), date(year, 3, 15), date(year, 5, 1), date(year, 8, 20),
        date(year, 10, 23), date(year, 11, 1), date(year, 12, 25), date(year, 12, 26),
    }
    a = year % 19
    b, c = divmod(year, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    ell = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * ell) // 451
    month = (h + ell - 7 * m + 114) // 31
    day = ((h + ell - 7 * m + 114) % 31) + 1
    easter = date(year, month, day)
    return fixed | {easter - timedelta(days=2), easter + timedelta(days=1), easter + timedelta(days=50)}


def is_workday(
    value: date,
    extra_workdays: set[date] | None = None,
    extra_holidays: set[date] | None = None,
) -> bool:
    if value in (extra_workdays or set()):
        return True
    if value in (extra_holidays or set()):
        return False
    return value.weekday() < 5 and value not in hungarian_holidays(value.year)


def add_calendar_months(value: date, months: int) -> date:
    """Shift by whole calendar months, clamping to the target month's last day."""
    month_index = value.year * 12 + value.month - 1 + months
    year, zero_based_month = divmod(month_index, 12)
    month = zero_based_month + 1
    day = min(value.day, monthrange(year, month)[1])
    return value.replace(year=year, month=month, day=day)


def calculate_schedule(
    rows: Iterable[Mapping[str, Any]],
    project_start: date,
    *,
    extra_workdays: set[date] | None = None,
    extra_holidays: set[date] | None = None,
) -> list[dict[str, Any]]:
    source = sorted(
        (dict(row) for row in rows),
        key=lambda row: (int(row.get("position") or 0), int(row.get("id") or 0)),
    )
    if not source:
        return []
    while not is_workday(project_start, extra_workdays, extra_holidays):
        project_start += timedelta(days=1)

    completed: dict[str, dict[str, Any]] = {}
    pending = source[:]
    while pending:
        progressed = False
        next_pending: list[dict[str, Any]] = []
        for row in pending:
            phase_name = str(row.get("phase_name") or "").strip()
            if not phase_name:
                raise ValueError("Minden ütemtervsorhoz fázisnév szükséges.")
            reference = str(row.get("reference_name") or PROJECT_START_LABEL).strip()
            manual = row.get("manual_start")
            if isinstance(manual, str) and manual:
                manual = date.fromisoformat(manual[:10])
            if manual:
                start = manual
            elif reference == PROJECT_START_LABEL:
                start = project_start
            elif reference in completed:
                start = completed[reference]["end_date"] + timedelta(days=1)
            else:
                next_pending.append(row)
                continue

            calc_type = row.get("calc_type") or "Naptári nap"
            if calc_type == "Munkanap":
                while not is_workday(start, extra_workdays, extra_holidays):
                    start += timedelta(days=1)
            duration = max(1, int(row.get("days") or 0) + int(row.get("delay") or 0))
            end = start
            if calc_type == "Hónap":
                end = add_calendar_months(start, duration)
            else:
                elapsed = 1
                while elapsed < duration:
                    end += timedelta(days=1)
                    if calc_type != "Munkanap" or is_workday(end, extra_workdays, extra_holidays):
                        elapsed += 1
            completed[phase_name] = {
                **row,
                "phase_name": phase_name,
                "reference_name": reference,
                "calc_type": calc_type,
                "start_date": start,
                "end_date": end,
            }
            progressed = True
        if not progressed:
            unresolved = ", ".join(str(row.get("phase_name") or "?") for row in next_pending)
            raise ValueError(f"Körkörös vagy hiányzó ütemtervi hivatkozás: {unresolved}")
        pending = next_pending
    return [completed[str(row["phase_name"]).strip()] for row in source]


def calculate_milestone_dates(
    milestones: Iterable[Mapping[str, Any]],
    calculated_schedule: Iterable[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Resolve fixed, schedule-linked and milestone-linked planned dates.

    Schedule references use the referenced phase's end date, matching the
    desktop application. Milestone references use the referenced milestone's
    planned/computed date; actual dates never move dependent milestones.
    """
    schedule_source = [dict(row) for row in calculated_schedule]
    source = [normalize_milestone(item) for item in milestones]
    schedule_end = {
        str(row.get("phase_name") or "").strip(): _as_date(row.get("end_date"))
        for row in schedule_source
    }
    schedule_end_by_id = {
        row_id: _as_date(row.get("end_date"))
        for row in schedule_source
        if (row_id := _as_positive_int(row.get("id"))) is not None
    }
    schedule_name_by_id = {
        row_id: str(row.get("phase_name") or "").strip()
        for row in schedule_source
        if (row_id := _as_positive_int(row.get("id"))) is not None
    }
    names: set[str] = set()
    for item in source:
        name = str(item.get("official_name") or "").strip()
        if not name:
            raise ValueError("Minden mérföldkőhöz név szükséges.")
        if name.casefold() in names:
            raise ValueError(f"Többször szereplő mérföldkőnév: {name}")
        names.add(name.casefold())

    completed: dict[str, date | None] = {}
    pending: list[dict[str, Any]] = []
    for item in source:
        date_type = str(item.get("date_type") or FIXED_DATE).strip()
        if date_type in {"", "Fix", FIXED_DATE}:
            completed[str(item["official_name"]).strip()] = _as_date(item.get("planned_date"))
        elif date_type in MILESTONE_DATE_TYPES:
            pending.append(item)
        else:
            raise ValueError(f"Ismeretlen mérföldkő-dátumtípus: {date_type}")

    while pending:
        progressed = False
        next_pending: list[dict[str, Any]] = []
        for item in pending:
            name = str(item["official_name"]).strip()
            date_type = str(item.get("date_type") or FIXED_DATE).strip()
            reference = str(item.get("reference_name") or "").strip()
            offset = int(item.get("offset_days") or 0)
            base: date | None
            if date_type == SCHEDULE_LINKED_DATE:
                schedule_row_id = _as_positive_int(item.get("schedule_row_id"))
                if schedule_row_id is not None and schedule_row_id in schedule_end_by_id:
                    base = schedule_end_by_id[schedule_row_id]
                    item["reference_name"] = schedule_name_by_id[schedule_row_id]
                else:
                    if schedule_row_id is not None:
                        item["schedule_row_id"] = None
                    if reference not in schedule_end:
                        raise ValueError(
                            f"A(z) „{name}” mérföldkő nem létező ütemtervi fázisra "
                            f"hivatkozik: {reference or 'nincs megadva'}."
                        )
                    base = schedule_end[reference]
            elif reference in completed:
                base = completed[reference]
            else:
                next_pending.append(item)
                continue
            completed[name] = base + timedelta(days=offset) if base else None
            progressed = True
        if not progressed:
            unresolved = ", ".join(str(item.get("name") or "?") for item in next_pending)
            raise ValueError(f"Körkörös vagy hiányzó mérföldkő-hivatkozás: {unresolved}")
        pending = next_pending

    result: list[dict[str, Any]] = []
    for item in source:
        resolved = completed[str(item["official_name"]).strip()]
        result.append({
            **item,
            "planned_date": resolved,
            "computed_planned_date": resolved,
        })
    return result


def build_schedule_workbook(
    project_name: str,
    project_start: date,
    rows: Iterable[Mapping[str, Any]],
    milestones: Iterable[Mapping[str, Any]],
) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    schedule_rows = [dict(item) for item in rows]
    milestone_rows = [dict(item) for item in milestones]
    dated = [
        value
        for value in [
            *(_as_date(item.get("start_date")) for item in schedule_rows),
            *(_as_date(item.get("end_date")) for item in schedule_rows),
            *(_as_date(item.get("planned_date")) for item in milestone_rows),
        ]
        if value
    ]
    start = min(dated, default=project_start) - timedelta(days=1)
    end = max(dated, default=project_start) + timedelta(days=20)
    if (end - start).days > 16000:
        raise ValueError("Az exportált napi idővonal túl hosszú az Excel oszlopkorlátjához.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ütemterv"
    sheet.freeze_panes = "E4"
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = f"{project_name} – ütemterv"
    sheet["A1"].font = Font(size=16, bold=True, color="173247")
    sheet["A2"] = "Projekt kezdete"
    sheet["B2"] = project_start
    sheet["B2"].number_format = "yyyy-mm-dd"
    headers = ["Szakasz / mérföldkő", "Kezdés", "Vége", "Nap"]
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(3, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="344F65")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    timeline: list[date] = []
    cursor = start
    while cursor <= end:
        timeline.append(cursor)
        cursor += timedelta(days=1)
    thin = Side(style="thin", color="C9D4DC")
    weekend_fill = PatternFill("solid", fgColor="FDE1DA")
    weekday_fill = PatternFill("solid", fgColor="EEF3F6")
    for index, day in enumerate(timeline, 5):
        cell = sheet.cell(3, index, day)
        cell.number_format = "dd"
        cell.alignment = Alignment(horizontal="center")
        cell.font = Font(size=9, bold=True, color="AD352E" if day.weekday() >= 5 else "173247")
        cell.fill = weekend_fill if day.weekday() >= 5 else weekday_fill
        cell.border = Border(left=thin, right=thin)
        sheet.column_dimensions[get_column_letter(index)].width = 3.8

    phase_colors = ["2E80BF", "1F9D61", "D68A16", "7255A8", "B84E4E"]
    row_no = 4
    for milestone in sorted(milestone_rows, key=lambda item: (int(item.get("phase_no") or 1), int(item.get("sequence_no") or 0), int(item.get("id") or 0))):
        planned = _as_date(milestone.get("planned_date"))
        sheet.cell(row_no, 1, f"◆ {milestone.get('name') or ''}")
        sheet.cell(row_no, 2, planned)
        sheet.cell(row_no, 3, planned)
        sheet.cell(row_no, 4, "Mérföldkő")
        color = phase_colors[(int(milestone.get("phase_no") or 1) - 1) % len(phase_colors)]
        if planned and start <= planned <= end:
            marker = sheet.cell(row_no, 5 + (planned - start).days, "◆")
            marker.font = Font(bold=True, color=color)
            marker.alignment = Alignment(horizontal="center")
        row_no += 1

    for index, item in enumerate(schedule_rows):
        row_start = _as_date(item.get("start_date"))
        row_end = _as_date(item.get("end_date"))
        sheet.cell(row_no, 1, item.get("phase_name") or "")
        sheet.cell(row_no, 2, row_start)
        sheet.cell(row_no, 3, row_end)
        sheet.cell(row_no, 4, max(1, (row_end - row_start).days + 1) if row_start and row_end else "")
        fill = PatternFill("solid", fgColor=phase_colors[index % len(phase_colors)])
        if row_start and row_end:
            for day_offset in range((row_end - row_start).days + 1):
                sheet.cell(row_no, 5 + (row_start - start).days + day_offset).fill = fill
        row_no += 1

    for column in (2, 3):
        for row in range(4, row_no):
            sheet.cell(row, column).number_format = "yyyy-mm-dd"
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions["C"].width = 13
    sheet.column_dimensions["D"].width = 12
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
