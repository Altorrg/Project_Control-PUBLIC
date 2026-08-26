from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from typing import Iterable

from .penzugy import parse_money


def contract_total(workflow: dict) -> Decimal:
    return parse_money(workflow.get("contract_fee")) + parse_money(workflow.get("copyright_fee"))


def offer_part_total(parts: Iterable[dict]) -> Decimal:
    return sum(
        (parse_money(part.get("amount")) for part in parts if not part.get("optional")),
        Decimal("0"),
    )


def validate_step_close(
    workflow: dict,
    offer_parts: list[dict],
    step: int,
    *,
    allow_offer_mismatch: bool = False,
) -> list[str]:
    warnings: list[str] = []
    if step == 1:
        if not str(workflow.get("company_name") or "").strip() or not str(workflow.get("discipline") or "").strip():
            raise ValueError("A cég és a szakág / szerep megadása kötelező.")
    elif step == 3 and not workflow.get("no_offer"):
        parts_total = offer_part_total(offer_parts)
        offer_total = parse_money(workflow.get("offer_price"))
        if parts_total and offer_total and parts_total != offer_total:
            message = (
                f"Az ajánlati ár ({offer_total:,.0f} Ft) nem egyezik a nem opcionális "
                f"részletek összegével ({parts_total:,.0f} Ft)."
            ).replace(",", " ")
            if not allow_offer_mismatch:
                raise ValueError(message)
            warnings.append(message)
    elif step == 5 and workflow.get("acceptance") == "Elfogadva":
        if not parse_money(workflow.get("contract_fee")) and not parse_money(workflow.get("copyright_fee")):
            raise ValueError(
                "Elfogadott ajánlatnál munkadíj vagy szerzői jogdíj szükséges a mátrix aktiválásához."
            )
    return warnings


def normalize_workflow_money(data: dict) -> dict:
    normalized = dict(data)
    normalized.pop("reopen_completed_payments", None)
    total = contract_total(normalized)
    normalized["contract_total"] = str(total.quantize(Decimal("0.01"))) if total else ""
    if normalized.get("no_offer"):
        normalized["offer_price"] = ""
    return normalized


def matrix_payment_reference(reference_id: str, column_code: str) -> str | None:
    parts = str(reference_id or "").upper().split("_")
    if len(parts) < 2 or parts[0] != "ALV" or not parts[1].isdigit():
        return None
    return f"KIA_WF_{int(parts[1])}_COL_{str(column_code or '').upper()}"


def build_subcontractor_workbook(
    folders: list[dict],
    workflows: list[dict],
) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ajánlatkérések"
    headers = [
        "Szakág", "Kód", "Alvállalkozó", "Előzetes kalkuláció", "Ajánlat státusza",
        "Ajánlati ár", "Elfogadás", "Szerződéses ár", "Módosított ár", "Aláírás",
        "Aktuális lépés", "Részletes bontás",
    ]
    thin = Side(style="thin", color="B8C5CE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="294A61")
    folder_fill = PatternFill("solid", fgColor="DCE6EC")
    signed_fill = PatternFill("solid", fgColor="E2F3EC")
    accepted_fill = PatternFill("solid", fgColor="FFF1D8")
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(1, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    grouped: dict[str, list[dict]] = defaultdict(list)
    for workflow in workflows:
        grouped[str(workflow.get("discipline") or "Egyéb")].append(workflow)
    folder_names = [str(item.get("name") or "Egyéb") for item in folders]
    folder_names.extend(name for name in grouped if name not in folder_names)

    row = 2
    for folder_name in folder_names:
        sheet.cell(row, 1, folder_name)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        folder_cell = sheet.cell(row, 1)
        folder_cell.font = Font(bold=True, color="17374B")
        folder_cell.fill = folder_fill
        folder_cell.border = border
        row += 1
        for workflow in sorted(grouped.get(folder_name, []), key=lambda item: (int(item.get("display_order") or 0), int(item.get("id") or 0))):
            details = "; ".join(
                f"{part.get('phase_name', '')}: {part.get('amount', '')}{' (opció)' if part.get('optional') else ''}"
                for part in workflow.get("offer_parts", [])
            )
            values = [
                folder_name,
                workflow.get("code", ""),
                workflow.get("company_name", ""),
                workflow.get("preliminary_estimate", ""),
                "Nem ad ajánlatot" if workflow.get("no_offer") else workflow.get("offer_status", ""),
                workflow.get("offer_price", ""),
                workflow.get("acceptance", ""),
                workflow.get("contract_total", ""),
                workflow.get("modified_total", ""),
                workflow.get("signature_status", ""),
                workflow.get("current_step", 1),
                details,
            ]
            fill = signed_fill if workflow.get("signature_status") == "Aláírva" else accepted_fill if workflow.get("acceptance") == "Elfogadva" else None
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column, value)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill
            row += 1

    widths = [24, 14, 34, 20, 20, 18, 18, 20, 20, 22, 15, 55]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(1, row - 1)}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
