from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


MAX_IMPORT_BYTES = 15 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
KNOWN_SHEETS = {"adatlap", "névadatlap", "osszegzes", "összegzés"}
STATUS_ALIASES = {
    "✓": "✓", "megvan": "✓", "elintezve": "✓", "elintézve": "✓",
    "○": "○", "folyamatban": "○", "elinditva": "○", "elindítva": "○",
    "!": "!", "nincs elinditva": "!", "nincs elindítva": "!",
    "⮾": "⮾", "elutasitva": "⮾", "elutasítva": "⮾",
    "h": "H", "hianypotlas": "H", "hiánypótlás": "H",
    "x": "X", "lejart": "X", "lejárt": "X",
    "ii": "II", "szuneteltetes": "II", "szüneteltetés": "II",
}


def validate_office_archive(content: bytes, filename: str) -> None:
    extension = Path(filename).suffix.lower()
    if extension not in {".xlsx", ".xlsm"}:
        raise ValueError("Csak XLSX vagy XLSM fajl importalhato.")
    if not content or len(content) > MAX_IMPORT_BYTES:
        raise ValueError("Az importfájl üres vagy túl nagy.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValueError("Az Office-csomag túl sok elemet tartalmaz.")
            total = 0
            for item in entries:
                normalized = item.filename.replace("\\", "/")
                if normalized.startswith("/") or "../" in normalized:
                    raise ValueError("Érvénytelen útvonal az Office-csomagban.")
                total += item.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("Az Office-csomag kibontott mérete túl nagy.")
    except BadZipFile as exc:
        raise ValueError("A feltöltött fájl nem érvényes Excel-munkafüzet.") from exc


def _text(value: Any, limit: int = 5000) -> str:
    return "" if value is None else str(value).strip()[:limit]


def _date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Hibás dátum: {text}")


def _status(value: Any) -> str:
    raw = _text(value, 80)
    if raw in STATUS_ALIASES:
        return STATUS_ALIASES[raw]
    folded = raw.casefold().replace(" / ", " ")
    for alias, code in STATUS_ALIASES.items():
        if alias.casefold() in folded:
            return code
    return "!"


def _header_row(sheet) -> int | None:
    for row_no in range(1, min(sheet.max_row, 15) + 1):
        values = [_text(sheet.cell(row_no, column).value, 100).casefold() for column in range(1, min(sheet.max_column, 12) + 1)]
        has_name = any(value in {"név", "nev", "partner", "partnernév"} for value in values)
        has_task = any("feladat" in value for value in values)
        if has_name and has_task:
            return row_no
    return 4 if sheet.max_row >= 5 and sheet.max_column >= 8 else None


def parse_engagement_workbook(content: bytes, filename: str) -> dict[str, Any]:
    validate_office_archive(content, filename)
    workbook = load_workbook(
        BytesIO(content),
        data_only=True,
        read_only=True,
        keep_links=False,
        keep_vba=False,
    )
    result: dict[str, Any] = {"tasks": [], "partners": [], "disciplines": [], "errors": [], "duplicates": []}
    seen_tasks: set[str] = set()
    seen_partners: set[str] = set()
    seen_assignments: set[tuple[str, str, str]] = set()

    try:
        if "Adatlap" in workbook.sheetnames:
            sheet = workbook["Adatlap"]
            for row in sheet.iter_rows(values_only=True):
                name = _text(row[0] if row else "", 255)
                if not name or name.casefold() in {"feladat", "feladat neve"}:
                    continue
                key = name.casefold()
                if key in seen_tasks:
                    continue
                seen_tasks.add(key)
                try:
                    weight = float(row[1] or 1) if len(row) > 1 else 1
                except (TypeError, ValueError):
                    weight = 1
                result["tasks"].append({
                    "name": name, "weight": weight,
                    "active": bool(row[2]) if len(row) > 2 and row[2] is not None else True,
                    "sort_order": int(row[3] or len(result["tasks"]) * 10) if len(row) > 3 else len(result["tasks"]) * 10,
                    "description": _text(row[4] if len(row) > 4 else "", 2000),
                })

        if "Névadatlap" in workbook.sheetnames:
            sheet = workbook["Névadatlap"]
            for row in sheet.iter_rows(values_only=True):
                name = _text(row[0] if row else "", 255)
                if not name or name.casefold() in {"név", "nev", "partnernév"}:
                    continue
                key = name.casefold()
                if key in seen_partners:
                    result["duplicates"].append({"type": "partner", "name": name})
                    continue
                seen_partners.add(key)
                result["partners"].append({
                    "name": name,
                    "address": _text(row[1] if len(row) > 1 else "", 500),
                    "contact_person": _text(row[2] if len(row) > 2 else "", 255),
                    "email": _text(row[3] if len(row) > 3 else "", 255),
                    "phone": _text(row[4] if len(row) > 4 else "", 80),
                    "org_type": _text(row[5] if len(row) > 5 else "", 120),
                })

        for sheet in workbook.worksheets:
            if sheet.title.casefold() in KNOWN_SHEETS or sheet.max_column < 8:
                continue
            header = _header_row(sheet)
            if header is None:
                result["errors"].append({"sheet": sheet.title, "message": "Nem található szakági fejléc."})
                continue
            discipline = {"name": _text(sheet.title, 120), "assignments": []}
            for row_no, row in enumerate(sheet.iter_rows(min_row=header + 1, values_only=True), start=header + 1):
                partner_name = _text(row[0] if row else "", 255)
                task_name = _text(row[2] if len(row) > 2 else "", 255)
                if not partner_name and not task_name:
                    continue
                if not partner_name or not task_name:
                    result["errors"].append({"sheet": sheet.title, "row": row_no, "message": "Partner vagy feladat hiányzik."})
                    continue
                duplicate_key = (discipline["name"].casefold(), partner_name.casefold(), task_name.casefold())
                if duplicate_key in seen_assignments:
                    result["duplicates"].append({"type": "assignment", "sheet": sheet.title, "row": row_no, "partner": partner_name, "task": task_name})
                    continue
                seen_assignments.add(duplicate_key)
                try:
                    assignment = {
                        "partner": {"name": partner_name, "address": _text(row[1] if len(row) > 1 else "", 500)},
                        "task_name": task_name,
                        "weight": float(row[3] or 1) if len(row) > 3 else 1,
                        "main_note": _text(row[4] if len(row) > 4 else "", 5000),
                        "last_date": _date(row[5]) if len(row) > 5 else None,
                        "status_code": _status(row[6] if len(row) > 6 else "!"),
                        "valid_until": _date(row[7]) if len(row) > 7 else None,
                        "rounds": [],
                    }
                    for index in range(8, len(row), 4):
                        values = list(row[index:index + 4]) + [None] * max(0, 4 - len(row[index:index + 4]))
                        if not any(value not in (None, "") for value in values):
                            continue
                        assignment["rounds"].append({
                            "sent_date": _date(values[0]), "sent_note": _text(values[1]),
                            "received_date": _date(values[2]), "received_note": _text(values[3]),
                            "status_code": assignment["status_code"],
                        })
                    discipline["assignments"].append(assignment)
                except (TypeError, ValueError) as exc:
                    result["errors"].append({"sheet": sheet.title, "row": row_no, "message": str(exc)})
            result["disciplines"].append(discipline)
    finally:
        workbook.close()
    result["summary"] = {
        "tasks": len(result["tasks"]), "partners": len(result["partners"]),
        "disciplines": len(result["disciplines"]),
        "assignments": sum(len(item["assignments"]) for item in result["disciplines"]),
        "rounds": sum(len(assignment["rounds"]) for item in result["disciplines"] for assignment in item["assignments"]),
        "errors": len(result["errors"]), "duplicates": len(result["duplicates"]),
    }
    return result


def build_engagement_excel(data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Adatlap"
    data_sheet.append(["Feladat", "Súly", "Aktív", "Sorrend", "Leírás"])
    for task in data.get("tasks", []):
        data_sheet.append([task.get("name"), task.get("weight"), task.get("active"), task.get("sort_order"), task.get("description")])
    names = workbook.create_sheet("Névadatlap")
    names.append(["Név", "Cím", "Ügyintéző", "E-mail", "Telefonszám", "Szervezet típus"])
    for partner in data.get("partners", []):
        names.append([partner.get("name"), partner.get("address"), partner.get("contact_person"), partner.get("email"), partner.get("phone"), partner.get("org_type")])
    summary = workbook.create_sheet("Összegzés")
    summary.append(["Szakág", "Ügyek", "Lezárt", "Figyelmeztetés", "Utolsó dátum", "Súlyozott készültség"])
    dashboard = data.get("dashboard", {})
    for discipline in dashboard.get("disciplines", []):
        summary.append([
            discipline.get("discipline_name"), discipline.get("total_count", 0), discipline.get("closed_count", 0),
            discipline.get("attention_count", 0), discipline.get("last_date", ""),
            float(discipline.get("progress_percent") or 0) / 100,
        ])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for discipline in data.get("disciplines", []):
        sheet = workbook.create_sheet(_text(discipline.get("name"), 31) or "Szakag")
        headers = ["Név", "Cím", "Feladat", "Súly", "Megjegyzés", "Utolsó dátum", "Státusz", "Érvényes"]
        for index in range(1, 11):
            headers.extend([f"{index}. KI dátum", f"{index}. KI megjegyzés", f"{index}. BE dátum", f"{index}. BE megjegyzés"])
        sheet.append(headers)
        for assignment in discipline.get("assignments", []):
            partner = assignment.get("partner") or {}
            task = assignment.get("task") or {}
            line = [
                partner.get("name") or assignment.get("temporary_partner_name") or "",
                assignment.get("address") or partner.get("address") or "", task.get("name") or assignment.get("task_name") or "",
                assignment.get("weight"), assignment.get("main_note"), assignment.get("last_date"),
                assignment.get("status_code"), assignment.get("valid_until"),
            ]
            rounds = assignment.get("rounds", [])[:10]
            for index in range(10):
                event = rounds[index] if index < len(rounds) else {}
                line.extend([event.get("sent_date", ""), event.get("sent_note", ""), event.get("received_date", ""), event.get("received_note", "")])
            sheet.append(line)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_engagement_pdf(data: dict[str, Any]) -> bytes:
    try:
        import reportlab
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError("A PDF-exporthoz a ReportLab csomag telepítése szükséges.") from exc
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    font_path = Path(reportlab.__file__).resolve().parent / "fonts" / "Vera.ttf"
    font_name = "ProjectControlVera"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
    _, height = A4
    y = height - 40
    pdf.setTitle("Engedélyezési dashboard riport")
    pdf.setFont(font_name, 14)
    pdf.drawString(40, y, "Engedélyek és hozzájárulások - dashboard riport")
    y -= 28
    dashboard = data.get("dashboard", {})
    pdf.setFont(font_name, 9)
    pdf.drawString(40, y, f"Ügyek: {dashboard.get('total', 0)} | Lezárt: {dashboard.get('closed', 0)} | Figyelmet igényel: {dashboard.get('attention', 0)}")
    y -= 22
    for item in dashboard.get("disciplines", []):
        line = f"{item.get('discipline_name', '')}: {item.get('total_count', 0)} ügy, {item.get('closed_count', 0)} lezárt, {item.get('progress_percent', 0)}%"
        pdf.drawString(40, y, line[:120])
        y -= 16
        if y < 50:
            pdf.showPage()
            pdf.setFont(font_name, 9)
            y = height - 40
    pdf.save()
    return output.getvalue()
