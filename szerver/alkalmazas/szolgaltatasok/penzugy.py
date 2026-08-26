from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Iterable, Mapping
import logging
import re


logger = logging.getLogger("project_control.finance_math")


def parse_money(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    normalized = re.sub(r"[^0-9,.-]", "", str(value))
    if "," in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    elif normalized.count(".") > 1 or re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", normalized):
        normalized = normalized.replace(".", "")
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return Decimal("0")


def finance_summary(
    contract_sections: Iterable[Mapping[str, Any]],
    project_contract_value: Any,
    daily_rates: Mapping[int, Any],
    daily_progress: Iterable[Mapping[str, Any]],
    workflows: Iterable[Mapping[str, Any]],
    payments: Iterable[Mapping[str, Any]],
    valid_derived_references: set[str] | None = None,
) -> dict[str, Any]:
    sections = list(contract_sections)
    if sections:
        budget = sum(
            (parse_money(row.get("modified_fee")) or parse_money(row.get("original_fee")))
            for row in sections
        )
    else:
        budget = parse_money(project_contract_value)

    current_internal = Decimal("0")
    estimated_internal = Decimal("0")
    discipline_totals: dict[str, dict[str, Decimal]] = {}
    for row in daily_progress:
        discipline = str(row.get("discipline_name") or "Ismeretlen")
        record_date = str(row.get("record_date") or "")
        year = int(record_date[:4]) if len(record_date) >= 4 and record_date[:4].isdigit() else 0
        cost = parse_money(row.get("days")) * parse_money(daily_rates.get(year, 0))
        percent = parse_money(row.get("actual_percent"))
        current_internal += cost
        stats = discipline_totals.setdefault(
            discipline,
            {"current": Decimal("0"), "actual_percent": Decimal("0")},
        )
        stats["current"] += cost
        stats["actual_percent"] += percent

    by_discipline: dict[str, dict[str, Decimal]] = {}
    for discipline, totals in discipline_totals.items():
        current = totals["current"]
        actual_percent = totals["actual_percent"]
        estimate = (
            current / (actual_percent / Decimal("100"))
            if Decimal("0") < actual_percent < Decimal("100")
            else current
        )
        estimated_internal += estimate
        by_discipline[discipline] = {"current": current, "estimated": estimate}
        logger.info(
            "finance_math discipline=%s current=%s actual_percent=%s estimated=%s",
            discipline,
            current,
            actual_percent,
            estimate,
        )

    paid_by_workflow: dict[int, Decimal] = {}
    for payment in payments:
        reference = str(payment.get("reference") or "").upper()
        if valid_derived_references is not None and re.fullmatch(r"KIA_WF_\d+_COL_(?:M|STAGE)_\d+", reference) and reference not in valid_derived_references:
            continue
        if payment.get("is_revenue") or payment.get("status") != "Pénzügyileg teljesítve":
            continue
        workflow_id = payment.get("workflow_id")
        if workflow_id is not None:
            paid_by_workflow[int(workflow_id)] = paid_by_workflow.get(int(workflow_id), Decimal("0")) + parse_money(payment.get("amount"))

    external_contracted = Decimal("0")
    external_paid = Decimal("0")
    subcontractors = []
    for workflow in workflows:
        if workflow.get("acceptance") != "Elfogadva":
            continue
        contracted = parse_money(workflow.get("modified_total")) or parse_money(workflow.get("contract_total"))
        paid = paid_by_workflow.get(int(workflow["id"]), Decimal("0"))
        external_contracted += contracted
        external_paid += paid
        subcontractors.append({**workflow, "contracted": contracted, "paid": paid, "remaining": max(Decimal("0"), contracted - paid)})

    estimated_total = estimated_internal + external_contracted
    logger.info(
        "finance_math summary budget=%s current_internal=%s estimated_internal=%s "
        "external_contracted=%s external_paid=%s estimated_total=%s profit=%s "
        "utilization_percent=%s",
        budget,
        current_internal,
        estimated_internal,
        external_contracted,
        external_paid,
        estimated_total,
        budget - estimated_total,
        (estimated_total / budget * Decimal("100")) if budget else Decimal("0"),
    )
    return {
        "budget": budget,
        "current_internal": current_internal,
        "estimated_internal": estimated_internal,
        "external_contracted": external_contracted,
        "external_paid": external_paid,
        "estimated_total_cost": estimated_total,
        "estimated_profit": budget - estimated_total,
        "utilization_percent": (estimated_total / budget * Decimal("100")) if budget else Decimal("0"),
        "by_discipline": by_discipline,
        "subcontractors": subcontractors,
    }


def build_finance_items(
    contract_sections: Iterable[Mapping[str, Any]],
    milestones: Iterable[Mapping[str, Any]],
    workflows: Iterable[Mapping[str, Any]],
    matrix_cells: Iterable[Mapping[str, Any]],
    payments: Iterable[Mapping[str, Any]],
    contract_stages: Iterable[Mapping[str, Any]] = (),
) -> list[dict[str, Any]]:
    """Builds the financial work queue from schedule milestones and matrix allocations."""
    section_rows = list(contract_sections)
    milestone_rows = list(milestones)
    workflow_rows = list(workflows)
    matrix_rows = list(matrix_cells)
    payment_rows = list(payments)
    stage_rows = list(contract_stages)
    sections = {int(row.get("section_no") or 1): row for row in section_rows}
    workflow_map = {
        int(row["id"]): row
        for row in workflow_rows
        if not row.get("is_deleted") and row.get("acceptance") == "Elfogadva"
    }
    payment_map = {str(row.get("reference") or "").upper(): row for row in payment_rows if row.get("reference")}
    items: list[dict[str, Any]] = []

    for milestone in milestone_rows:
        milestone_id = int(milestone["id"])
        section = sections.get(int(milestone.get("phase_no") or 1))
        if section is None:
            continue
        section_fee = parse_money(section.get("modified_fee")) or parse_money(section.get("original_fee"))
        fee_base = section_fee * parse_money(section.get("fee_percent")) / Decimal("100")
        copyright_base = section_fee * parse_money(section.get("copyright_percent")) / Decimal("100")
        expected = (
            fee_base * parse_money(milestone.get("fee_percent")) / Decimal("100")
            + copyright_base * parse_money(milestone.get("copyright_percent")) / Decimal("100")
        )
        logger.info(
            "finance_math revenue milestone_id=%s phase_no=%s section_fee=%s "
            "section_fee_percent=%s fee_base=%s milestone_fee_percent=%s "
            "section_copyright_percent=%s copyright_base=%s "
            "milestone_copyright_percent=%s expected=%s rounding=database_numeric_18_2",
            milestone_id,
            milestone.get("phase_no") or 1,
            section_fee,
            section.get("fee_percent") or 0,
            fee_base,
            milestone.get("fee_percent") or 0,
            section.get("copyright_percent") or 0,
            copyright_base,
            milestone.get("copyright_percent") or 0,
            expected,
        )
        reference = f"BEV_MS_{milestone_id}"
        payment = payment_map.get(reference)
        items.append({
            "reference": reference,
            "is_revenue": True,
            "workflow_id": None,
            "source_milestone_id": milestone_id,
            "title": str(milestone.get("name") or reference),
            "partner": "Főszerződés",
            "expected_date": milestone.get("actual_date") or milestone.get("planned_date"),
            "expected_amount": expected,
            "payment": payment,
        })

    cost_totals: dict[tuple[int, str], Decimal] = {}
    for cell in matrix_rows:
        column_code = str(cell.get("column_code") or "").strip().upper()
        workflow_match = re.match(r"ALV_(\d+)_", str(cell.get("reference_id") or ""), re.IGNORECASE)
        if not re.fullmatch(r"(?:M|STAGE)_\d+", column_code) or not workflow_match:
            continue
        key = (int(workflow_match.group(1)), column_code)
        cost_totals[key] = cost_totals.get(key, Decimal("0")) + parse_money(cell.get("amount"))
        logger.info(
            "finance_math cost_partial workflow_id=%s column=%s reference=%s "
            "cell_percent=%s cell_amount=%s running_column_total=%s",
            key[0],
            key[1],
            cell.get("reference_id"),
            cell.get("percent"),
            parse_money(cell.get("amount")),
            cost_totals[key],
        )

    milestone_map = {int(row["id"]): row for row in milestone_rows}
    stage_map = {int(row["id"]): row for row in stage_rows}
    for (workflow_id, column_code), expected in sorted(cost_totals.items(), key=lambda item: (item[0][1], item[0][0])):
        if expected <= 0:
            continue
        workflow = workflow_map.get(workflow_id)
        if workflow is None:
            continue
        source_id = int(column_code.rsplit("_", 1)[1])
        milestone = milestone_map.get(source_id, {}) if column_code.startswith("M_") else {}
        stage = stage_map.get(source_id, {}) if column_code.startswith("STAGE_") else {}
        if not milestone and not stage:
            continue
        reference = f"KIA_WF_{workflow_id}_COL_{column_code}"
        payment = payment_map.get(reference)
        logger.info(
            "finance_math cost_item workflow_id=%s column=%s expected=%s "
            "vat_mode=%s vat_conversion=not_applied currency=%s",
            workflow_id,
            column_code,
            expected,
            workflow.get("vat_type") or "Nettó",
            workflow.get("currency") or "HUF",
        )
        items.append({
            "reference": reference,
            "is_revenue": False,
            "workflow_id": workflow_id,
            "source_milestone_id": source_id if milestone else None,
            "source_stage_id": source_id if stage else None,
            "title": str(milestone.get("name") or stage.get("name") or column_code),
            "partner": str(workflow.get("company_name") or f"Alvállalkozó {workflow_id}"),
            "expected_date": milestone.get("actual_date") or milestone.get("planned_date") or stage.get("stage_date"),
            "expected_amount": expected,
            "payment": payment,
        })

    result = sorted(items, key=lambda row: (str(row.get("expected_date") or "9999-12-31"), not row["is_revenue"], row["partner"]))
    logger.info(
        "finance_math queue section_count=%s milestone_count=%s workflow_count=%s "
        "matrix_cell_count=%s payment_count=%s item_count=%s revenue_total=%s cost_total=%s",
        len(section_rows),
        len(milestone_rows),
        len(workflow_rows),
        len(matrix_rows),
        len(payment_rows),
        len(result),
        sum((parse_money(item["expected_amount"]) for item in result if item["is_revenue"]), Decimal("0")),
        sum((parse_money(item["expected_amount"]) for item in result if not item["is_revenue"]), Decimal("0")),
    )
    return result


def company_cash_flow(
    year: int,
    projects: Iterable[Mapping[str, Any]],
    payments: Iterable[Mapping[str, Any]],
    daily_rates: Iterable[Mapping[str, Any]],
    daily_progress: Iterable[Mapping[str, Any]],
) -> dict[str, Any]:
    rate_map = {
        int(row["project_id"]): parse_money(row.get("daily_rate"))
        for row in daily_rates
        if int(row.get("year") or 0) == year
    }
    project_map: dict[int, dict[str, Any]] = {}
    for project in projects:
        project_id = int(project["id"])
        project_map[project_id] = {
            "project_id": project_id,
            "project_name": str(project.get("name") or f"{project_id}. projekt"),
            "contract_value": parse_money(project.get("contract_value")),
            "revenue": [Decimal("0") for _ in range(12)],
            "external_cost": [Decimal("0") for _ in range(12)],
            "internal_cost": [Decimal("0") for _ in range(12)],
        }

    def month_index(value: Any) -> int | None:
        text = str(value or "")
        if len(text) < 7 or not text[:4].isdigit() or not text[5:7].isdigit() or int(text[:4]) != year:
            return None
        month = int(text[5:7])
        return month - 1 if 1 <= month <= 12 else None

    for payment in payments:
        project = project_map.get(int(payment.get("project_id") or 0))
        month = month_index(payment.get("payment_date"))
        if project is None or month is None:
            continue
        key = "revenue" if payment.get("is_revenue") else "external_cost"
        project[key][month] += parse_money(payment.get("amount"))

    for progress in daily_progress:
        project_id = int(progress.get("project_id") or 0)
        project = project_map.get(project_id)
        month = month_index(progress.get("record_date"))
        if project is None or month is None:
            continue
        project["internal_cost"][month] += parse_money(progress.get("days")) * rate_map.get(project_id, Decimal("0"))

    monthly_revenue = [Decimal("0") for _ in range(12)]
    monthly_external = [Decimal("0") for _ in range(12)]
    monthly_internal = [Decimal("0") for _ in range(12)]
    project_rows = []
    for project in project_map.values():
        project["net"] = [
            project["revenue"][month] - project["external_cost"][month] - project["internal_cost"][month]
            for month in range(12)
        ]
        project["annual_net"] = sum(project["net"], Decimal("0"))
        project["annual_revenue"] = sum(project["revenue"], Decimal("0"))
        project["annual_cost"] = sum(project["external_cost"], Decimal("0")) + sum(project["internal_cost"], Decimal("0"))
        for month in range(12):
            monthly_revenue[month] += project["revenue"][month]
            monthly_external[month] += project["external_cost"][month]
            monthly_internal[month] += project["internal_cost"][month]
        project_rows.append(project)

    monthly_cost = [monthly_external[index] + monthly_internal[index] for index in range(12)]
    monthly_net = [monthly_revenue[index] - monthly_cost[index] for index in range(12)]
    cumulative = []
    running = Decimal("0")
    for value in monthly_net:
        running += value
        cumulative.append(running)

    return {
        "year": year,
        "contract_total": sum((row["contract_value"] for row in project_rows), Decimal("0")),
        "revenue_total": sum(monthly_revenue, Decimal("0")),
        "cost_total": sum(monthly_cost, Decimal("0")),
        "profit_total": sum(monthly_net, Decimal("0")),
        "monthly_revenue": monthly_revenue,
        "monthly_external_cost": monthly_external,
        "monthly_internal_cost": monthly_internal,
        "monthly_cost": monthly_cost,
        "monthly_net": monthly_net,
        "cumulative": cumulative,
        "projects": sorted(project_rows, key=lambda row: row["project_name"].casefold()),
    }


def labor_trends(
    daily_progress: Iterable[Mapping[str, Any]],
    daily_rates: Mapping[int, Any],
) -> dict[str, Any]:
    rows = list(daily_progress)
    months = sorted({str(row.get("record_date") or "")[:7] for row in rows if len(str(row.get("record_date") or "")) >= 7})
    disciplines = sorted({str(row.get("discipline_name") or "Ismeretlen") for row in rows}, key=str.casefold)
    monthly: dict[tuple[str, str], dict[str, Decimal]] = {}
    for row in rows:
        month = str(row.get("record_date") or "")[:7]
        if month not in months:
            continue
        discipline = str(row.get("discipline_name") or "Ismeretlen")
        values = monthly.setdefault((discipline, month), {"days": Decimal("0"), "planned": Decimal("0"), "actual": Decimal("0"), "cost": Decimal("0")})
        days = parse_money(row.get("days"))
        record_date = str(row.get("record_date") or "")
        year = int(record_date[:4]) if len(record_date) >= 4 and record_date[:4].isdigit() else 0
        values["days"] += days
        values["planned"] += parse_money(row.get("planned_percent"))
        values["actual"] += parse_money(row.get("actual_percent"))
        values["cost"] += days * parse_money(daily_rates.get(year, 0))

    series = []
    for discipline in disciplines:
        running_days = Decimal("0")
        cumulative_days = []
        monthly_days = []
        monthly_cost = []
        planned = []
        actual = []
        for month in months:
            values = monthly.get((discipline, month), {"days": Decimal("0"), "planned": Decimal("0"), "actual": Decimal("0"), "cost": Decimal("0")})
            running_days += values["days"]
            monthly_days.append(values["days"])
            cumulative_days.append(running_days)
            monthly_cost.append(values["cost"])
            planned.append(values["planned"])
            actual.append(values["actual"])
        series.append({
            "discipline": discipline,
            "monthly_days": monthly_days,
            "cumulative_days": cumulative_days,
            "monthly_cost": monthly_cost,
            "planned_percent": planned,
            "actual_percent": actual,
        })

    overall_planned = []
    overall_actual = []
    for month in months:
        active = [monthly[(discipline, month)] for discipline in disciplines if (discipline, month) in monthly and (monthly[(discipline, month)]["planned"] or monthly[(discipline, month)]["actual"])]
        divisor = Decimal(len(active)) if active else Decimal("1")
        overall_planned.append(sum((item["planned"] for item in active), Decimal("0")) / divisor)
        overall_actual.append(sum((item["actual"] for item in active), Decimal("0")) / divisor)
    return {"months": months, "series": series, "overall_planned": overall_planned, "overall_actual": overall_actual}


def finance_allocations(
    by_discipline: Mapping[str, Mapping[str, Any]],
    subcontractors: Iterable[Mapping[str, Any]],
    mappings: Iterable[Mapping[str, Any]],
    discipline_names: Iterable[str],
) -> dict[str, Any]:
    mapping_by_workflow = {int(row["workflow_id"]): str(row.get("discipline_name") or "") for row in mappings}
    names = list(dict.fromkeys([*discipline_names, *by_discipline.keys()]))
    rows: dict[str, dict[str, Any]] = {
        name: {
            "discipline": name,
            "internal_current": parse_money(by_discipline.get(name, {}).get("current")),
            "internal_estimated": parse_money(by_discipline.get(name, {}).get("estimated")),
            "external_paid": Decimal("0"),
            "external_contracted": Decimal("0"),
            "subcontractors": [],
        }
        for name in names
    }
    unassigned = {"discipline": "Nincs szakághoz rendelve", "internal_current": Decimal("0"), "internal_estimated": Decimal("0"), "external_paid": Decimal("0"), "external_contracted": Decimal("0"), "subcontractors": []}
    for item in subcontractors:
        workflow_id = int(item["id"])
        target_name = mapping_by_workflow.get(workflow_id, "")
        target = rows.get(target_name) if target_name else None
        if target is None:
            target = unassigned
        target["external_paid"] += parse_money(item.get("paid"))
        target["external_contracted"] += parse_money(item.get("contracted"))
        target["subcontractors"].append({
            "id": workflow_id,
            "company_name": str(item.get("company_name") or ""),
            "contracted": parse_money(item.get("contracted")),
            "paid": parse_money(item.get("paid")),
        })
    return {"disciplines": list(rows.values()), "unassigned": unassigned, "mappings": list(mappings)}


def reconciliation_report(finance_items: Iterable[Mapping[str, Any]], payments: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    items = list(finance_items)
    payment_rows = list(payments)
    source_map = {str(item.get("reference") or "").upper(): item for item in items}
    reference_counts: dict[str, int] = {}
    issues = []
    for payment in payment_rows:
        reference = str(payment.get("reference") or "").strip().upper()
        if reference:
            reference_counts[reference] = reference_counts.get(reference, 0) + 1
        if re.fullmatch(r"(?:BEV_MS_\d+|KIA_WF_\d+_COL_(?:M|STAGE)_\d+)", reference) and reference not in source_map:
            issues.append({"type": "orphan", "severity": "error", "payment_id": payment.get("id"), "reference": reference, "message": "Az eredeti mérföldkő vagy mátrixcella már nem létezik."})
            continue
        source = source_map.get(reference)
        if source and str(payment.get("status") or "") == "Pénzügyileg teljesítve" and abs(parse_money(payment.get("amount")) - parse_money(source.get("expected_amount"))) >= Decimal("0.01"):
            issues.append({"type": "amount_mismatch", "severity": "warning", "payment_id": payment.get("id"), "reference": reference, "message": "A tényleges és a forrás szerinti összeg eltér."})
    for reference, count in reference_counts.items():
        if reference and count > 1:
            issues.append({"type": "duplicate", "severity": "error", "reference": reference, "message": f"Ugyanahhoz a forráshoz {count} pénzügyi rekord tartozik."})
    for item in items:
        if not item.get("expected_date"):
            issues.append({"type": "missing_date", "severity": "warning", "reference": item.get("reference"), "message": "A pénzügyi forráshoz nincs várható dátum."})
    return {
        "issues": issues,
        "error_count": sum(1 for issue in issues if issue["severity"] == "error"),
        "warning_count": sum(1 for issue in issues if issue["severity"] == "warning"),
    }


def build_project_finance_workbook(project_name: str, summary: Mapping[str, Any]) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Vezetői dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.append(["Project Control - Pénzügyi vezetői riport"])
    dashboard["A1"].font = Font(size=16, bold=True, color="243746")
    dashboard.append(["Projekt", project_name])
    dashboard.append([])
    for label, key in (
        ("Szerződéses keret", "budget"),
        ("Várható belső költség", "estimated_internal"),
        ("Szerződött alvállalkozók", "external_contracted"),
        ("Várható eredmény", "estimated_profit"),
    ):
        dashboard.append([label, summary.get(key, 0)])
    for cell in dashboard[4]:
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.font = Font(color="FFFFFF", bold=True)

    allocation_sheet = workbook.create_sheet("Költségallokáció")
    allocation_sheet.append(["Szakág", "Belső jelenlegi", "Belső várható", "Külső kifizetett", "Külső szerződött"])
    allocation_rows = list(summary.get("allocations", {}).get("disciplines", []))
    unassigned = summary.get("allocations", {}).get("unassigned", {})
    if unassigned.get("subcontractors"):
        allocation_rows.append(unassigned)
    for row in allocation_rows:
        allocation_sheet.append([row.get("discipline"), row.get("internal_current", 0), row.get("internal_estimated", 0), row.get("external_paid", 0), row.get("external_contracted", 0)])

    labor_sheet = workbook.create_sheet("Munkaórák")
    labor_sheet.append(["Szakág", "Hónap", "Havi nap", "Kumulált nap", "Tervezett %", "Tényleges %"])
    trends = summary.get("labor_trends", {})
    for series in trends.get("series", []):
        for index, month in enumerate(trends.get("months", [])):
            labor_sheet.append([series.get("discipline"), month, series.get("monthly_days", [])[index], series.get("cumulative_days", [])[index], series.get("planned_percent", [])[index], series.get("actual_percent", [])[index]])

    payment_sheet = workbook.create_sheet("Pénzügyi folyamatok")
    payment_sheet.append(["Típus", "Partner", "Forrás", "Várható dátum", "Várható összeg", "Státusz", "TIG", "Számla", "Teljesítés", "Tényleges összeg"])
    for item in summary.get("finance_items", []):
        payment = item.get("payment") or {}
        payment_sheet.append(["Bevétel" if item.get("is_revenue") else "Kiadás", item.get("partner"), item.get("title"), item.get("expected_date"), item.get("expected_amount", 0), payment.get("status", "Várható"), payment.get("tig_date"), payment.get("invoice_date"), payment.get("payment_date"), payment.get("amount", 0)])

    for cell in dashboard["B"][3:7]:
        cell.number_format = '#,##0 "Ft"'
    for row in allocation_sheet.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = '#,##0 "Ft"'
    for row in labor_sheet.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '0.0 "nap"'
    for row in labor_sheet.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = '0.0"%"'
    for column in ("E", "J"):
        for cell in payment_sheet[column][1:]:
            cell.number_format = '#,##0 "Ft"'

    for sheet in workbook.worksheets:
        if sheet.max_row:
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="243746")
                cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
