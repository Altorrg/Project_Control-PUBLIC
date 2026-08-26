from __future__ import annotations

import asyncio
from io import BytesIO
import logging
import re
from typing import Annotated

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from ..hitelesites import AccessToken, AuthenticatedUser, CsrfProtected, get_gateway
from ..adatmodellek.penzugy import FinanceMapping, PaymentCreate
from ..hibak import ConflictError
from ..jogosultsagok import require_permission
from ..szolgaltatasok.penzugy import (
    build_finance_items,
    build_project_finance_workbook,
    company_cash_flow,
    finance_allocations,
    finance_summary,
    labor_trends,
    reconciliation_report,
)
from ..supabase_kapcsolat import SupabaseGateway


router = APIRouter(prefix="/finance", tags=["finance"])
DERIVED_PAYMENT_REFERENCE = re.compile(r"^(?:BEV_MS_\d+|KIA_WF_\d+_COL_(?:M|STAGE)_\d+)$", re.IGNORECASE)
logger = logging.getLogger("project_control.finance_math")


async def _project_finance_items(project_id: int, token: str, gateway: SupabaseGateway, payments: list[dict] | None = None) -> list[dict]:
    sections, milestones, workflows, matrix_cells, stages = await asyncio.gather(
        gateway.select("project_contract_sections", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("project_milestones", token, filters={"project_id": f"eq.{project_id}"}, order="planned_date.asc,sequence_no.asc,id.asc"),
        gateway.select("subcontractor_workflows", token, filters={"project_id": f"eq.{project_id}", "is_deleted": "eq.false"}),
        gateway.select("contract_matrix", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("contract_stages", token, filters={"project_id": f"eq.{project_id}"}, order="sort_order.asc,id.asc"),
    )
    return build_finance_items(sections.data, milestones.data, workflows.data, matrix_cells.data, payments or [], stages.data)


async def _validate_payment_source(project_id: int, data: dict, token: str, gateway: SupabaseGateway) -> dict:
    raw_reference = str(data.get("reference") or "").strip()
    reference = raw_reference.upper()
    if not DERIVED_PAYMENT_REFERENCE.fullmatch(reference):
        data["reference"] = raw_reference
        return data
    data["reference"] = reference
    sources = await _project_finance_items(project_id, token, gateway)
    source = next((item for item in sources if item["reference"].upper() == reference), None)
    if source is None:
        raise ConflictError("A pénzügyi tétel forrása már nem létezik az ütemtervben vagy a kifizetési mátrixban.")
    data["workflow_id"] = source.get("workflow_id")
    data["is_revenue"] = bool(source.get("is_revenue"))
    return data


async def _validate_unique_payment_reference(project_id: int, data: dict, token: str, gateway: SupabaseGateway, payment_id: int | None = None) -> None:
    reference = str(data.get("reference") or "").strip()
    if not reference:
        return
    existing = await gateway.select(
        "payments",
        token,
        filters={"project_id": f"eq.{project_id}", "reference": f"eq.{reference}"},
    )
    if any(int(row.get("id") or 0) != int(payment_id or 0) for row in existing.data):
        raise ConflictError("Ehhez a forráshivatkozáshoz már tartozik pénzügyi tétel a projektben.")


async def _validate_finance_mapping(project_id: int, data: dict, token: str, gateway: SupabaseGateway) -> None:
    workflow_id = int(data.get("workflow_id") or 0)
    discipline_name = str(data.get("discipline_name") or "").strip()
    workflow, disciplines = await asyncio.gather(
        gateway.select(
            "subcontractor_workflows",
            token,
            filters={"id": f"eq.{workflow_id}", "project_id": f"eq.{project_id}", "is_deleted": "eq.false"},
            limit=1,
        ),
        gateway.select(
            "project_disciplines",
            token,
            filters={"project_id": f"eq.{project_id}", "visible": "eq.true"},
        ),
    )
    if not workflow.data:
        raise ConflictError("Az alvállalkozói szerződés nem tartozik ehhez a projekthez vagy már törölték.")
    if discipline_name not in {str(row.get("name") or "").strip() for row in disciplines.data}:
        raise ConflictError("A kiválasztott szakág nem tartozik az aktuális projekthez.")


async def _company_cash_flow_data(year: int, token: str, gateway: SupabaseGateway) -> dict:
    start = f"{year}-01-01"
    end = f"{year}-12-31"
    projects = await gateway.select("projects", token, filters={"is_deleted": "eq.false"}, order="name.asc")
    payments = await gateway.select("payments", token, filters={"payment_date": f"gte.{start}", "and": f"(payment_date.lte.{end})"})
    rates = await gateway.select("project_daily_rates", token, filters={"year": f"eq.{year}"})
    progress = await gateway.select("daily_progress", token, filters={"record_date": f"gte.{start}", "and": f"(record_date.lte.{end})"})
    sections = await gateway.select("project_contract_sections", token)
    milestones = await gateway.select("project_milestones", token)
    workflows = await gateway.select("subcontractor_workflows", token, filters={"is_deleted": "eq.false"})
    matrix_cells = await gateway.select("contract_matrix", token)
    stages = await gateway.select("contract_stages", token)

    def by_project(rows: list[dict], project_id: int) -> list[dict]:
        return [row for row in rows if int(row.get("project_id") or 0) == project_id]

    valid_derived_references: set[tuple[int, str]] = set()
    for project in projects.data:
        project_id = int(project["id"])
        items = build_finance_items(
            by_project(sections.data, project_id),
            by_project(milestones.data, project_id),
            by_project(workflows.data, project_id),
            by_project(matrix_cells.data, project_id),
            [],
            by_project(stages.data, project_id),
        )
        valid_derived_references.update((project_id, str(item["reference"]).upper()) for item in items)
    effective_payments = [
        payment
        for payment in payments.data
        if not DERIVED_PAYMENT_REFERENCE.fullmatch(str(payment.get("reference") or "").upper())
        or (int(payment.get("project_id") or 0), str(payment.get("reference") or "").upper()) in valid_derived_references
    ]
    return company_cash_flow(year, projects.data, effective_payments, rates.data, progress.data)


async def _project_summary_data(project_id: int, token: str, gateway: SupabaseGateway) -> dict:
    projects, sections, rates, progress, workflows, payments, milestones, matrix_cells, stages, mappings, disciplines = await asyncio.gather(
        gateway.select("projects", token, filters={"id": f"eq.{project_id}"}, limit=1),
        gateway.select("project_contract_sections", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("project_daily_rates", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("daily_progress", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("subcontractor_workflows", token, filters={"project_id": f"eq.{project_id}", "is_deleted": "eq.false"}),
        gateway.select("payments", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("project_milestones", token, filters={"project_id": f"eq.{project_id}"}, order="planned_date.asc,sequence_no.asc,id.asc"),
        gateway.select("contract_matrix", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("contract_stages", token, filters={"project_id": f"eq.{project_id}"}, order="sort_order.asc,id.asc"),
        gateway.select("finance_mappings", token, filters={"project_id": f"eq.{project_id}"}),
        gateway.select("project_disciplines", token, filters={"project_id": f"eq.{project_id}", "visible": "eq.true"}, order="sort_order.asc,name.asc"),
    )
    daily_rates = {int(row["year"]): row["daily_rate"] for row in rates.data}
    project = projects.data[0] if projects.data else {}
    finance_items = build_finance_items(sections.data, milestones.data, workflows.data, matrix_cells.data, payments.data, stages.data)
    valid_references = {str(item["reference"]).upper() for item in finance_items}
    result = finance_summary(sections.data, project.get("contract_value", 0), daily_rates, progress.data, workflows.data, payments.data, valid_references)
    result["project_name"] = project.get("name", "Projekt")
    result["finance_items"] = finance_items
    result["labor_trends"] = labor_trends(progress.data, daily_rates)
    result["allocations"] = finance_allocations(result["by_discipline"], result["subcontractors"], mappings.data, [str(row.get("name") or "") for row in disciplines.data])
    result["reconciliation"] = reconciliation_report(finance_items, payments.data)
    logger.info(
        "finance_math project_summary project_id=%s sections=%s milestones=%s "
        "workflows=%s matrix_cells=%s payments=%s finance_items=%s "
        "budget=%s estimated_total=%s estimated_profit=%s",
        project_id,
        len(sections.data),
        len(milestones.data),
        len(workflows.data),
        len(matrix_cells.data),
        len(payments.data),
        len(finance_items),
        result.get("budget"),
        result.get("estimated_total_cost"),
        result.get("estimated_profit"),
    )
    return result


@router.get("/company-cash-flow")
async def company_summary(
    token: AccessToken,
    _: AuthenticatedUser,
    __=Depends(require_permission("arbevetel")),
    gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None,
    year: int = Query(ge=2000, le=2200),
) -> dict:
    return await _company_cash_flow_data(year, token, gateway)


@router.get("/company-cash-flow/export")
async def export_company_summary(
    token: AccessToken,
    _: AuthenticatedUser,
    __=Depends(require_permission("arbevetel")),
    gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None,
    year: int = Query(ge=2000, le=2200),
) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = await _company_cash_flow_data(year, token, gateway)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Cash-flow_{year}"
    sheet.append(["Project Control - céges cash-flow", year])
    sheet["A1"].font = Font(size=15, bold=True)
    sheet.append([])
    sheet.append(["Főszerződések", data["contract_total"]])
    sheet.append(["Éves bevétel", data["revenue_total"]])
    sheet.append(["Éves költség", data["cost_total"]])
    sheet.append(["Éves nettó egyenleg", data["profit_total"]])
    sheet.append([])
    headers = ["Hónap", "Bevétel", "Alvállalkozói költség", "Belső ráfordítás", "Nettó", "Göngyölített"]
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.font = Font(color="FFFFFF", bold=True)
    month_names = ["Január", "Február", "Március", "Április", "Május", "Június", "Július", "Augusztus", "Szeptember", "Október", "November", "December"]
    for index, month in enumerate(month_names):
        sheet.append([month, data["monthly_revenue"][index], data["monthly_external_cost"][index], data["monthly_internal_cost"][index], data["monthly_net"][index], data["cumulative"][index]])
    sheet.append([])
    sheet.append(["Projekt", "Éves bevétel", "Éves költség", "Éves nettó"])
    for project in data["projects"]:
        sheet.append([project["project_name"], project["annual_revenue"], project["annual_cost"], project["annual_net"]])
    for column in ("B", "C", "D", "E", "F"):
        for cell in sheet[column]:
            if isinstance(cell.value, (int, float)) or cell.value.__class__.__name__ == "Decimal":
                cell.number_format = '#,##0 "Ft"'
    sheet.column_dimensions["A"].width = 30
    for column in ("B", "C", "D", "E", "F"):
        sheet.column_dimensions[column].width = 20
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="project_control_cash_flow_{year}.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/{project_id}/summary")
async def summary(project_id: int, token: AccessToken, _: AuthenticatedUser, gateway: Annotated[SupabaseGateway, Depends(get_gateway)]) -> dict:
    return await _project_summary_data(project_id, token, gateway)


@router.get("/{project_id}/export.xlsx")
async def export_project_finance(project_id: int, token: AccessToken, _: AuthenticatedUser, gateway: Annotated[SupabaseGateway, Depends(get_gateway)]) -> StreamingResponse:
    data = await _project_summary_data(project_id, token, gateway)
    output = build_project_finance_workbook(str(data.get("project_name") or "Projekt"), data)
    headers = {"Content-Disposition": f'attachment; filename="project_control_penzugy_{project_id}.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/{project_id}/payments")
async def list_payments(project_id: int, token: AccessToken, _: AuthenticatedUser, gateway: Annotated[SupabaseGateway, Depends(get_gateway)]) -> list[dict]:
    result = await gateway.select("payments", token, filters={"project_id": f"eq.{project_id}"}, order="payment_date.desc,id.desc")
    return result.data


@router.post("/{project_id}/payments")
async def create_payment(project_id: int, payload: PaymentCreate, token: AccessToken, _: CsrfProtected, __=Depends(require_permission("finance_write")), gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None) -> dict:
    data = payload.model_dump(mode="json")
    data["project_id"] = project_id
    data = await _validate_payment_source(project_id, data, token, gateway)
    logger.info(
        "finance_math payment_create project_id=%s reference=%s amount=%s "
        "status=%s is_revenue=%s",
        project_id,
        data.get("reference"),
        data.get("amount"),
        data.get("status"),
        data.get("is_revenue"),
    )
    await _validate_unique_payment_reference(project_id, data, token, gateway)
    result = await gateway.insert("payments", token, data)
    return result[0]


@router.put("/{project_id}/payments/{payment_id}")
async def update_payment(project_id: int, payment_id: int, payload: PaymentCreate, token: AccessToken, _: CsrfProtected, __=Depends(require_permission("finance_write")), gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None) -> dict:
    data = payload.model_dump(mode="json")
    data["project_id"] = project_id
    data = await _validate_payment_source(project_id, data, token, gateway)
    logger.info(
        "finance_math payment_update project_id=%s payment_id=%s reference=%s "
        "amount=%s status=%s is_revenue=%s",
        project_id,
        payment_id,
        data.get("reference"),
        data.get("amount"),
        data.get("status"),
        data.get("is_revenue"),
    )
    await _validate_unique_payment_reference(project_id, data, token, gateway, payment_id)
    result = await gateway.update("payments", token, data, {"id": f"eq.{payment_id}", "project_id": f"eq.{project_id}"})
    return result[0]


@router.delete("/{project_id}/payments/{payment_id}")
async def delete_payment(project_id: int, payment_id: int, token: AccessToken, _: CsrfProtected, __=Depends(require_permission("finance_write")), gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None) -> dict[str, bool]:
    await gateway.delete("payments", token, {"id": f"eq.{payment_id}", "project_id": f"eq.{project_id}"})
    return {"deleted": True}


@router.put("/{project_id}/mappings")
async def map_workflow(project_id: int, payload: FinanceMapping, token: AccessToken, _: CsrfProtected, __=Depends(require_permission("finance_write")), gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None) -> dict:
    data = payload.model_dump(mode="json")
    data["project_id"] = project_id
    data["discipline_name"] = str(data["discipline_name"]).strip()
    await _validate_finance_mapping(project_id, data, token, gateway)
    result = await gateway.upsert("finance_mappings", token, data, on_conflict="project_id,workflow_id")
    return result[0]


@router.delete("/{project_id}/mappings/{workflow_id}")
async def unmap_workflow(project_id: int, workflow_id: int, token: AccessToken, _: CsrfProtected, __=Depends(require_permission("finance_write")), gateway: Annotated[SupabaseGateway, Depends(get_gateway)] = None) -> dict[str, bool]:
    await gateway.delete("finance_mappings", token, {"project_id": f"eq.{project_id}", "workflow_id": f"eq.{workflow_id}"})
    return {"deleted": True}
